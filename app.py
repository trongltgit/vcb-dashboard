"""
VCB Dashboard – full app
- Plan pie: đọc KH+TH từ file KE_HOACH (logic Excel gốc)
- Bar TOP/BOT DS+LN MBNT + TTQT
- PDF, AI (GROQ_MODEL=openai/gpt-oss-20b)
"""
import base64
import gc
import io
import os
import re
import unicodedata as ud
from datetime import date, datetime

import numpy as np
import pandas as pd
from dotenv import load_dotenv
from flask import (
    Flask, flash, jsonify, redirect, render_template,
    request, send_file, session, url_for,
)
from werkzeug.utils import secure_filename

from utils.processing import prepare_all, build_customer_df, filt, build, kpi_pair
from utils.ai_bot import generate_professional_report, chat_with_data

load_dotenv()

app = Flask(__name__)
app.secret_key = os.environ.get("SECRET_KEY", "vcb-dashboard-dev-key")
app.config["MAX_CONTENT_LENGTH"] = 80 * 1024 * 1024

ENABLE_CHARTS = os.environ.get("ENABLE_CHARTS", "1") == "1"
PLAN_YEAR = int(os.environ.get("PLAN_YEAR", "2026"))
_STORE = {}


def _sid():
    if "sid" not in session:
        session["sid"] = os.urandom(8).hex()
    return session["sid"]


# ─── IO ───────────────────────────────────────────────

def read_df(fs):
    if fs is None or not fs.filename:
        return None
    name = (secure_filename(fs.filename) or "f").lower()
    fs.stream.seek(0)
    if name.endswith(".csv"):
        try:
            df = pd.read_csv(fs.stream, encoding="utf-8-sig", low_memory=False)
        except UnicodeDecodeError:
            fs.stream.seek(0)
            df = pd.read_csv(fs.stream, encoding="cp1258", low_memory=False)
    elif name.endswith((".xlsx", ".xlsm")):
        df = pd.read_excel(fs.stream, engine="openpyxl")
    elif name.endswith(".xls"):
        df = pd.read_excel(fs.stream, engine="xlrd")
    else:
        raise ValueError(f"Dinh dang khong ho tro: {name}")
    df.columns = [str(c).strip() for c in df.columns]
    return df


def read_external_text(fs) -> str:
    if fs is None or not fs.filename:
        return ""
    name = fs.filename.lower()
    data = fs.read()
    if name.endswith((".txt", ".md")):
        return data.decode("utf-8", errors="ignore")
    if name.endswith(".pdf"):
        try:
            from pypdf import PdfReader
            return "\n".join(p.extract_text() or "" for p in PdfReader(io.BytesIO(data)).pages)
        except Exception:
            return ""
    if name.endswith(".docx"):
        try:
            from docx import Document
            return "\n".join(p.text for p in Document(io.BytesIO(data)).paragraphs)
        except Exception:
            return ""
    return ""


def usd(v):
    a = abs(v)
    if a >= 1e9: return f"{a/1e9:,.2f} ty USD"
    if a >= 1e6: return f"{a/1e6:,.1f} trieu USD"
    if a >= 1e3: return f"{a/1e3:,.0f} nghin USD"
    return f"{a:,.0f} USD"


def pc(a, b):
    return f"{a/b*100:.0f}%" if b > 0 else "0%"


def safe_html(df, n=30):
    if df is None or df.empty:
        return ""
    try:
        d = df.drop(columns=["Key"], errors="ignore").head(n).astype(str)
        return d.to_html(classes="table table-sm", index=False, escape=True, border=0)
    except Exception:
        return ""


# ─── Plan: EXACT Excel KE_HOACH logic ─────────────────

def _norm(s):
    return " ".join(str(s).replace("\xa0", " ").split()).lower()


def _code_of(s):
    s = ud.normalize("NFD", str(s).replace("đ", "d").replace("Đ", "D"))
    s = "".join(ch for ch in s if ud.category(ch) != "Mn")
    return "".join(w if w.isdigit() else w[0] for w in re.findall(r"[A-Za-z]+|\d+", s)).lower()


def _num(x):
    try:
        v = float(x)
        return 0.0 if pd.isna(v) else v
    except Exception:
        return 0.0


def build_plan_grid_excel(plan_raw, phong, d1, d2, year=PLAN_YEAR):
    """
    Port dung logic Python Excel (ROWS + V + th_cum + cells).
    Doc KH va TH tu file KE_HOACH upload — khong lay TH tu raw.
    """
    if plan_raw is None or plan_raw.empty:
        return None, None, None, None

    plan = plan_raw.copy()
    while plan.shape[1] < 15:
        plan[plan.shape[1]] = 0
    if len(plan) < 186:
        extra = pd.DataFrame([[None] * plan.shape[1]] * (186 - len(plan)))
        plan = pd.concat([plan, extra], ignore_index=True)

    def V(r, col):
        # r 1-index, col 0-index
        try:
            return _num(plan.iloc[r - 1, col])
        except Exception:
            return 0.0

    rooms_n = [_norm(plan.iloc[5 + 3 * i, 0]) for i in range(15)]  # row 6 Excel = index 5

    d1 = pd.Timestamp(d1)
    d2 = pd.Timestamp(d2)
    m1, m2 = d1.month, d2.month
    if d1.year < year:
        m1 = 1
    if d2.year > year:
        m2 = 12

    QN = (d2.month - 1) // 3 + 1
    # Excel: Q_COL = 1 + d2.month  (0-index col of month T{m2})
    # Cot: 0=ten, 1=KH/TH, 2=T1 ... 13=T12, 14=Ca Nam
    Q_COL = 1 + m2          # T{m2}
    Y_COL = 14              # Ca Nam
    Q_LABEL = f"Q{QN}/{d2.year}"
    Y_LABEL = f"{d2.year}"

    METRICS = [
        dict(name="DS MBNT", unit="tr.USD"),
        dict(name="LN MBNT", unit="ty VND"),
        dict(name="DS TTQT-TTTM", unit="tr.USD"),
    ]
    # Row 1-index Excel (giong code goc)
    ROWS = [
        dict(k0=6,   cn=54,  tsc=60,  tot=63),
        dict(k0=69,  cn=117, tsc=123, tot=126),
        dict(k0=132, cn=177, tsc=180, tot=183),
    ]

    def is_cum(k):
        R = ROWS[k]
        for r in (R["cn"] + 1, R["tot"] + 1):
            xs = [V(r, 1 + m) for m in range(1, m2 + 1)]
            if m2 >= 3 and xs[-1] > 0:
                up = sum(1 for i in range(1, m2) if xs[i] > xs[i - 1])
                return all(xs[i] >= xs[i - 1] - 1e-9 for i in range(1, m2)) and up >= max(2, int(0.6 * (m2 - 1)))
        return False

    CUM = [is_cum(k) for k in range(3)]

    def th_cum(k, r_th):
        if CUM[k]:
            return V(r_th, 1 + m2) - (V(r_th, m1) if m1 > 1 else 0.0)
        return float(sum(V(r_th, 1 + m) for m in range(m1, m2 + 1)))

    key = _norm(phong or "")
    codes = [_code_of(x) for x in rooms_n]
    if key and key in rooms_n:
        i_room = rooms_n.index(key)
    elif key and key.replace(" ", "") in codes:
        i_room = codes.index(key.replace(" ", ""))
    else:
        i_room = None
    phong_mode = i_room is not None

    def cells(k, kh_col):
        R = ROWS[k]
        r_phong = (R["k0"] + 3 * i_room) if phong_mode else R["tot"]
        out = []
        for r_kh in (r_phong, R["cn"], R["tsc"]):
            th = th_cum(k, r_kh + 1)   # dong TH = KH + 1
            kh = V(r_kh, kh_col)
            out.append((th, kh, th / kh if kh > 0 else 0.0))
        return out

    grid = [cells(k, Q_COL) + cells(k, Y_COL) for k in range(3)]
    return grid, METRICS, Q_LABEL, Y_LABEL


def render_plan_chart(plan_raw, phong, d_from, d_to):
    grid, metrics, q_label, y_label = build_plan_grid_excel(plan_raw, phong, d_from, d_to)
    if grid is None:
        return "", []
    try:
        import matplotlib
        matplotlib.use("Agg")
        import matplotlib.pyplot as plt
        from utils.charts import chart_plan_completion

        fig = chart_plan_completion(
            grid, metrics,
            loc=phong or "VCB HCM",
            d1=pd.Timestamp(d_from),
            d2=pd.Timestamp(d_to),
            q_label=q_label,
            y_label=y_label,
            bg="#14532D",
        )
        buf = io.BytesIO()
        fig.savefig(buf, format="png", bbox_inches="tight", dpi=100, facecolor="#14532D")
        plt.close(fig)
        buf.seek(0)
        b64 = base64.b64encode(buf.read()).decode("ascii")

        names = ["KH phong", "KH CN giao", "KH TSC giao"]
        summary = []
        for mi, m in enumerate(metrics):
            for ci, cn in enumerate(names):
                th, kh, pct = grid[mi][ci]
                summary.append(f"{m['name']} | {cn} Q: TH={th:,.2f} KH={kh:,.2f} ({pct*100:.1f}%)")
            for ci, cn in enumerate(names):
                th, kh, pct = grid[mi][ci + 3]
                summary.append(f"{m['name']} | {cn} Y: TH={th:,.2f} KH={kh:,.2f} ({pct*100:.1f}%)")
        return b64, summary
    except Exception as e:
        print(f"Plan chart failed: {e}")
        return "", []


# ─── Charts ───────────────────────────────────────────

def generate_charts_safe(TDS, BDS, TLN, BLN, TQ, BQ, q_filt, df):
    charts = {k: "" for k in [
        "mbnt_top_ds", "mbnt_bot_ds", "mbnt_top_ln", "mbnt_bot_ln",
        "ttqt_top", "ttqt_bot", "ttqt_xk_top", "ttqt_xk_bot",
        "ttqt_nk_top", "ttqt_nk_bot", "scatter", "plan_pie",
    ]}
    try:
        import matplotlib
        matplotlib.use("Agg")
        import matplotlib.pyplot as plt
        from utils.charts import chart_bar, scatter_ttqt_mbnt

        def to_b64(fig):
            buf = io.BytesIO()
            fig.savefig(buf, format="png", bbox_inches="tight", dpi=72)
            plt.close(fig)
            buf.seek(0)
            return base64.b64encode(buf.read()).decode("ascii")

        # Core charts – luon ve
        charts["mbnt_top_ds"] = to_b64(chart_bar(TDS, "TOP 20% DS MBNT", "USD", "#0EA5E9", True))
        charts["mbnt_bot_ds"] = to_b64(chart_bar(BDS, "BOT 10% DS MBNT", "USD", "#F97316", False))
        charts["mbnt_top_ln"] = to_b64(chart_bar(TLN, "TOP 20% LN MBNT", "VND", "#22C55E", True))
        charts["mbnt_bot_ln"] = to_b64(chart_bar(BLN, "BOT 10% LN MBNT", "VND", "#EF4444", False))
        charts["ttqt_top"] = to_b64(chart_bar(TQ, "TOP 20% DS TTQT", "USD", "#8B5CF6", True))
        charts["ttqt_bot"] = to_b64(chart_bar(BQ, "BOT 10% DS TTQT", "USD", "#EC4899", False))

        if ENABLE_CHARTS:
            cq_xk = build(q_filt, ["xk"]).rename(columns={"XK": "DS"})
            cq_nk = build(q_filt, ["nk"]).rename(columns={"NK": "DS"})
            TXK, BXK = kpi_pair(cq_xk, "DS")
            TNK, BNK = kpi_pair(cq_nk, "DS")
            charts["ttqt_xk_top"] = to_b64(chart_bar(TXK, "TOP XK", "USD", "#06B6D4", True))
            charts["ttqt_xk_bot"] = to_b64(chart_bar(BXK, "BOT XK", "USD", "#F59E0B", False))
            charts["ttqt_nk_top"] = to_b64(chart_bar(TNK, "TOP NK", "USD", "#10B981", True))
            charts["ttqt_nk_bot"] = to_b64(chart_bar(BNK, "BOT NK", "USD", "#F43F5E", False))
            charts["scatter"] = to_b64(scatter_ttqt_mbnt(df))
            del TXK, BXK, TNK, BNK, cq_xk, cq_nk
        gc.collect()
    except Exception as e:
        print(f"Charts failed: {e}")
    return charts


# ─── Report ───────────────────────────────────────────

def build_report(df, dall, qall, mm, tt, phong, d_from, d_to):
    def phan_tich(X, O, Y, Yv, tx_, ty_):
        L = []
        d = df[df[X] > 0].copy()
        if d.empty:
            return [f"Khong co cong ty co {tx_}."], {}
        TX, TY, TV = d[X].sum(), d[Y].sum(), d[Yv].sum()
        r = d[Y] / d[X]
        L.append(f"Co {len(d):,} cty co {tx_}, tong {usd(TX)}. "
                 f"{int((d[Y]>0).sum()):,} cty co {ty_}, tong {usd(TY)}.")
        L.append(f"Chuyen doi: 100 USD {tx_} -> {TY/TX*100:.0f} USD {ty_}.")
        lab = ["Khong GD", "Duoi 30%", "30-80%", "80-120%", "Tren 120%"]
        d["Nhom"] = np.select([d[Y] == 0, r < 0.3, r < 0.8, r <= 1.2], lab[:4], default=lab[4])
        for k in lab:
            s = d[d.Nhom == k]
            if len(s):
                L.append(f"  - {k}: {len(s):,} cty, {pc(s[X].sum(), TX)} {tx_}.")
        lo = d[d.Nhom.isin(lab[:2])]
        L.append(f"Nhom <30%: {len(lo):,} cty, {pc(lo[X].sum(), TX)} {tx_}, "
                 f"{pc(lo[Y].sum(), TY)} {ty_}.")
        L.append(f"Giai ngan vay NT: {pc(TV, TY)} {ty_} ({usd(TV)}).")
        d["Net"] = (d[X] - d[O]).clip(lower=0)
        d["Tn"] = (d.Net - d[Y]).clip(lower=0)
        low = d[d[Y] < 0.5 * d[X]].copy()
        low["Nn"] = np.select([low[O] >= 0.7 * low[X], low[Y] == 0], ["A", "B"], default="C")
        tn = low[low.Nn.isin(["B", "C"])].Tn.sum()
        L.append(f"Tiem nang B+C: {usd(tn)}.")
        return L, {"conv": TY / TX if TX else 0, "tn": tn}

    LX, sx = phan_tich("XK", "NK", "BUY", "BUY_VAY", "XK", "BUY")
    LN, sn = phan_tich("NK", "XK", "SELL", "SELL_VAY", "NK", "SELL")
    R = [
        "BAO CAO PHAN TICH DS TTQT - DS MBNT",
        f"Phong: {phong or 'VCB HCM'}. Ky: {d_from:%d/%m/%Y} - {d_to:%d/%m/%Y}.",
        f"MBNT {len(dall):,} dong; TTQT {len(qall):,}; {len(df):,} DN.",
        "", "PHAN 1. XK VA BUY",
    ] + LX + ["", "PHAN 2. NK VA SELL"] + LN + ["", "PHAN 3. KET LUAN"]
    if sx and sn:
        R.append(f"Chuyen doi XK->BUY {sx['conv']*100:.0f}%, NK->SELL {sn['conv']*100:.0f}%.")
        R.append("Giu chan tot hon o chieu mua NT." if sn["conv"] > sx["conv"]
                 else "Giu chan tot hon o chieu ban NT.")
        R.append(f"Tiem nang lon hon o {'XK' if sx['tn'] >= sn['tn'] else 'NK'}.")
    return "\n".join(R)


# ─── PDF ──────────────────────────────────────────────

def _pdf_safe(s):
    if not s:
        return ""
    for a, b in [("–", "-"), ("—", "-"), ("→", "->"), ("đ", "d"), ("Đ", "D"),
                 ("“", '"'), ("”", '"'), ("…", "...")]:
        s = s.replace(a, b)
    return s.encode("latin-1", errors="replace").decode("latin-1")


def build_pdf(data):
    from fpdf import FPDF

    class PDF(FPDF):
        def header(self):
            self.set_font("Helvetica", "B", 10)
            self.set_text_color(30, 58, 138)
            self.cell(0, 6, "VCB Analytics - MBNT & TTQT", new_x="LMARGIN", new_y="NEXT")
            self.line(10, self.get_y(), 200, self.get_y())
            self.ln(3)

        def footer(self):
            self.set_y(-12)
            self.set_font("Helvetica", "I", 8)
            self.set_text_color(120, 120, 120)
            self.cell(0, 8, f"Page {self.page_no()}/{{nb}}", align="C")

    pdf = PDF()
    pdf.alias_nb_pages()
    pdf.set_auto_page_break(True, 15)
    pdf.set_margins(12, 15, 12)
    pdf.add_page()
    w = pdf.w - 24

    def line(t, h=5, bold=False):
        t = _pdf_safe(str(t)).strip()
        if not t:
            pdf.ln(2)
            return
        pdf.set_font("Helvetica", "B" if bold else "", 11 if bold else 9)
        pdf.multi_cell(w, h, t)

    line("Bao cao phan tich MBNT & TTQT", 8, True)
    line(f"Phong: {data.get('phong') or 'VCB HCM'} | {data.get('d_from')} -> {data.get('d_to')}")
    pdf.ln(2)
    m = data.get("metrics") or {}
    line("Tong quan", True)
    line(f"MBNT DS: {m.get('mbnt_ds',0):,.0f} USD | LN: {m.get('mbnt_ln',0):,.0f} VND")
    line(f"MBNT GD: {m.get('mbnt_gd',0):,} | CIF: {m.get('mbnt_cif',0):,}")
    line(f"TTQT DS: {m.get('ttqt_ds',0):,.0f} | XK: {m.get('ttqt_xk',0):,.0f} | NK: {m.get('ttqt_nk',0):,.0f}")
    pdf.ln(2)
    for s in (data.get("plan_summary") or []):
        line(s, 4)
    pdf.ln(2)
    line("Bao cao chi tiet", True)
    for s in _pdf_safe(data.get("report_text") or "").splitlines():
        line(s, 4)

    for key, title in [
        ("plan_pie", "Hoan thanh ke hoach"),
        ("mbnt_top_ds", "TOP DS MBNT"), ("mbnt_bot_ds", "BOT DS MBNT"),
        ("mbnt_top_ln", "TOP LN MBNT"), ("mbnt_bot_ln", "BOT LN MBNT"),
        ("ttqt_top", "TOP TTQT"), ("ttqt_bot", "BOT TTQT"),
        ("ttqt_xk_top", "TOP XK"), ("ttqt_nk_top", "TOP NK"),
        ("scatter", "Scatter"),
    ]:
        b64 = (data.get("charts") or {}).get(key) or ""
        if not b64:
            continue
        try:
            pdf.add_page()
            line(title, True)
            pdf.image(io.BytesIO(base64.b64decode(b64)), x=12, w=w)
        except Exception as e:
            print(f"PDF img {key}: {e}")

    out = io.BytesIO()
    pdf.output(out)
    return out.getvalue()


# ─── Routes ───────────────────────────────────────────

@app.errorhandler(413)
def too_large(_e):
    flash("File > 80MB.", "error")
    return redirect(url_for("index")), 413


@app.route("/")
def index():
    return render_template("index.html", has_data=bool(_STORE.get(_sid())), today=date.today().isoformat())


@app.route("/run", methods=["POST"])
def run_analysis():
    mbnt_f = request.files.get("mbnt")
    ttqt_f = request.files.get("ttqt")
    plan_f = request.files.get("plan")
    ext_f = request.files.get("external")
    phong = (request.form.get("phong") or "").strip()
    d_from_s = request.form.get("d_from") or "2026-01-01"
    d_to_s = request.form.get("d_to") or "2026-12-31"

    if not mbnt_f or not ttqt_f or not mbnt_f.filename or not ttqt_f.filename:
        flash("Can upload MBNT + TTQT.", "error")
        return redirect(url_for("index"))

    try:
        d_from = datetime.strptime(d_from_s, "%Y-%m-%d").date()
        d_to = datetime.strptime(d_to_s, "%Y-%m-%d").date()
    except ValueError:
        flash("Ngay sai.", "error")
        return redirect(url_for("index"))
    if d_from > d_to:
        flash("Ngay bat dau > ket thuc.", "error")
        return redirect(url_for("index"))

    try:
        mbnt_raw = read_df(mbnt_f)
        ttqt_raw = read_df(ttqt_f)
        plan_raw = read_df(plan_f) if plan_f and plan_f.filename else None
        if mbnt_raw is None or ttqt_raw is None:
            raise ValueError("Khong doc duoc file.")

        dall, qall, SA, SQ = prepare_all(mbnt_raw, ttqt_raw)
        del mbnt_raw, ttqt_raw
        gc.collect()

        df, mm, tt, note = build_customer_df(dall, qall, phong, d_from, d_to)
        ext_text = read_external_text(ext_f)
        report_text = build_report(df, dall, qall, mm, tt, phong, d_from, d_to)

        d_filt, _ = filt(dall, phong, d_from, d_to)
        q_filt, _ = filt(qall, phong, d_from, d_to)

        cif_mbnt = build(d_filt, ["ds", "ln"])
        TDS, BDS = kpi_pair(cif_mbnt, "DS")
        TLN, BLN = kpi_pair(cif_mbnt, "LN")
        cq = build(q_filt, ["xk", "nk", "xnk", "ds"])
        TQ, BQ = kpi_pair(cq, "DS")

        metrics = {
            "mbnt_ds": float(d_filt.ds.sum()) if len(d_filt) else 0,
            "mbnt_ln": float(d_filt.ln.sum()) if len(d_filt) else 0,
            "mbnt_gd": int(len(d_filt)),
            "mbnt_cif": int(d_filt.loc[d_filt.cif != "", "cif"].nunique()) if len(d_filt) else 0,
            "ttqt_ds": float(q_filt.ds.sum()) if len(q_filt) else 0,
            "ttqt_xk": float(q_filt.xk.sum()) if len(q_filt) else 0,
            "ttqt_nk": float(q_filt.nk.sum()) if len(q_filt) else 0,
        }

        # Plan pie TU FILE KE_HOACH (KH+TH trong file)
        plan_pie_b64, plan_summary = "", []
        if plan_raw is not None:
            plan_pie_b64, plan_summary = render_plan_chart(plan_raw, phong, d_from, d_to)

        charts = generate_charts_safe(TDS, BDS, TLN, BLN, TQ, BQ, q_filt, df)
        charts["plan_pie"] = plan_pie_b64

        groups = {}
        try:
            for side, X, Y in [("XK -> BUY", "XK", "BUY"), ("NK -> SELL", "NK", "SELL")]:
                if X not in df.columns:
                    continue
                dside = df[pd.to_numeric(df[X], errors="coerce").fillna(0) > 0].copy()
                if dside.empty:
                    continue
                dside[X] = pd.to_numeric(dside[X], errors="coerce").fillna(0)
                dside[Y] = pd.to_numeric(dside[Y], errors="coerce").fillna(0)
                r = dside[Y] / dside[X].replace(0, np.nan)
                lab = ["1. Khong GD", "2. <30%", "3. 30-80%", "4. 80-120%", "5. >120%"]
                dside["Nhom"] = np.select([dside[Y] == 0, r < 0.3, r < 0.8, r <= 1.2], lab[:4], default=lab[4])
                summary = dside.groupby("Nhom", observed=True).agg(
                    So_cty=("Ten", "size"), Tong_X=(X, "sum"), Tong_Y=(Y, "sum")
                ).reset_index()
                summary["Y/X_%"] = np.where(summary["Tong_X"] > 0,
                    (summary["Tong_Y"] / summary["Tong_X"] * 100).round(1), 0)
                groups[side] = summary.to_dict(orient="records")
        except Exception as ge:
            print(f"groups: {ge}")

        phong_list = sorted(set(
            str(x).strip() for x in list(dall.ph.dropna().unique()) + list(qall.ph.dropna().unique())
            if str(x).strip() and str(x).strip().lower() not in ("nan", "none", "")
        ))

        _STORE[_sid()] = {
            "note": note, "report_text": report_text, "ext_text": ext_text,
            "charts": charts, "groups": groups, "metrics": metrics,
            "plan_summary": plan_summary, "plan_rows": [],
            "df_csv": df.to_csv(index=True).encode("utf-8-sig"),
            "tds_html": safe_html(TDS), "bds_html": safe_html(BDS),
            "phong": phong, "phong_list": phong_list,
            "d_from": d_from.isoformat(), "d_to": d_to.isoformat(),
            "ai_report": None, "chat_history": [],
        }
        msg = note
        if plan_raw is None:
            msg += " | Chua upload KE_HOACH"
        elif not plan_pie_b64:
            msg += " | KE_HOACH parse loi (kiem tra layout dong 6/69/132)"
        flash(msg, "success")

        del dall, qall, df, mm, tt, d_filt, q_filt
        del cif_mbnt, TDS, BDS, TLN, BLN, cq, TQ, BQ, SA, SQ, plan_raw
        gc.collect()
    except Exception as e:
        flash(f"Loi: {e}", "error")
        gc.collect()
        return redirect(url_for("index"))
    return redirect(url_for("dashboard"))


@app.route("/dashboard")
def dashboard():
    data = _STORE.get(_sid())
    if not data:
        flash("Chua co du lieu.", "error")
        return redirect(url_for("index"))
    return render_template("dashboard.html", data=data)


@app.route("/export/csv")
def export_csv():
    data = _STORE.get(_sid())
    if not data:
        return redirect(url_for("index"))
    return send_file(io.BytesIO(data["df_csv"]), mimetype="text/csv",
                     as_attachment=True, download_name="customer_analysis.csv")


@app.route("/export/report")
def export_report():
    data = _STORE.get(_sid())
    if not data:
        return redirect(url_for("index"))
    return send_file(io.BytesIO(data["report_text"].encode("utf-8")), mimetype="text/plain",
                     as_attachment=True, download_name="bao_cao.txt")


@app.route("/export/pdf")
def export_pdf():
    data = _STORE.get(_sid())
    if not data:
        return redirect(url_for("index"))
    try:
        return send_file(io.BytesIO(build_pdf(data)), mimetype="application/pdf",
                         as_attachment=True, download_name="vcb_bao_cao.pdf")
    except Exception as e:
        flash(f"PDF loi: {e}", "error")
        return redirect(url_for("dashboard"))


@app.route("/ai/generate", methods=["POST"])
def ai_generate():
    data = _STORE.get(_sid())
    if not data:
        return jsonify({"error": "Chua co du lieu"}), 400
    ctx = data.get("report_text", "")
    if data.get("ext_text"):
        ctx += "\n\n=== EXTERNAL ===\n" + data["ext_text"][:8000]
    report = generate_professional_report(ctx)
    data["ai_report"] = report
    return jsonify({"report": report})


@app.route("/ai/chat", methods=["POST"])
def ai_chat():
    data = _STORE.get(_sid())
    if not data:
        return jsonify({"error": "Chua co du lieu"}), 400
    msg = (request.json or {}).get("message", "").strip()
    if not msg:
        return jsonify({"error": "Empty"}), 400
    hist = data.get("chat_history") or []
    hist.append({"role": "user", "content": msg})
    ctx = data.get("report_text", "")
    if data.get("ext_text"):
        ctx += "\n\n=== EXTERNAL ===\n" + data["ext_text"][:8000]
    ans = chat_with_data(hist, ctx)
    hist.append({"role": "assistant", "content": ans})
    data["chat_history"] = hist
    return jsonify({"answer": ans, "history": hist})


@app.route("/reset", methods=["POST"])
def reset():
    _STORE.pop(_sid(), None)
    gc.collect()
    return redirect(url_for("index"))


if __name__ == "__main__":
    app.run(host="0.0.0.0", port=int(os.environ.get("PORT", 5000)), debug=True)
